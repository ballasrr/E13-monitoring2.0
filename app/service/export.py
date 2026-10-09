"""Выгрузка реестра в Excel, CSV и JSON.

Строки собираются один раз в методах *_table, а отрисовка в .xlsx или
.csv — отдельная функция. Так колонки описаны в одном месте, и формат
файла не влияет на то, что в нём лежит.

Excel — основной формат: в нём нет вопроса кодировки, который на
Windows отравляет любой CSV с кириллицей. CSV оставлен для случаев,
когда файл надо отдать скрипту, а не человеку.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.domain import compliance as rules
from app.domain import documents as docs
from app.models.enums import CHARGER_STATUS_RU, STATION_STATUS_RU, AlertLevel
from app.repositories.station import StationRepository
from app.service.compliance import station_to_facts


@dataclass(slots=True)
class Table:
    """Готовый к отрисовке лист: имя, заголовки, строки."""

    name: str
    headers: list[str]
    rows: list[list[Any]]


# Пары «поле модели — заголовок». Порядок колонок задаётся здесь и
# больше нигде: добавили поле в модель — добавили строку сюда.
STATION_COLUMNS: tuple[tuple[str, str], ...] = (
    ("code", "Код"),
    ("name", "Название"),
    ("status", "Стадия"),
    ("region", "Район"),
    ("address", "Адрес"),
    ("site_type", "Тип площадки"),
    ("lat", "Широта"),
    ("lng", "Долгота"),
    ("land_status", "Статус земли"),
    ("landlord", "Арендодатель"),
    ("opened_at", "Дата открытия"),
    ("grid_company", "Сетевая компания"),
    ("tp_number", "Номер ТП"),
    ("transformer_kva", "Трансформатор, кВА"),
    ("allocated_kw", "Мощность, кВт"),
    ("voltage", "Напряжение"),
    ("connection_contract", "Договор ТП"),
    ("connection_date", "Дата ТП"),
    ("meter_number", "Номер счётчика"),
    ("energy_supplier", "Энергосбыт"),
    ("has_canopy", "Навес"),
    ("has_lighting", "Освещение"),
    ("has_cctv", "Видеонаблюдение"),
    ("has_internet", "Интернет"),
    ("internet_type", "Тип интернета"),
    ("has_fence", "Ограждение"),
    ("has_wc", "Туалет"),
    ("has_cafe", "Кафе"),
    ("has_signage", "Навигация"),
    ("accessible", "Доступность МГН"),
    ("parking_spots", "Парковочных мест"),
    ("surface_type", "Покрытие"),
    ("has_operator", "Оператор на месте"),
    ("staff_count", "Сотрудников"),
    ("work_schedule", "Режим работы"),
    ("responsible_name", "Ответственный"),
    ("responsible_phone", "Телефон"),
    ("service_company", "Обслуживающая компания"),
    ("notes", "Примечания"),
)

CHARGER_COLUMNS: tuple[tuple[str, str], ...] = (
    ("name", "Название"),
    ("vendor", "Производитель"),
    ("model", "Модель"),
    ("serial", "Серийный номер"),
    ("ocpp_id", "OCPP ID"),
    ("status", "Состояние"),
    ("power_kw", "Мощность, кВт"),
    ("current_type", "Тип тока"),
    ("connectors", "Коннекторы"),
    ("connector_count", "Портов"),
    ("tariff_rub", "Тариф, руб/кВт*ч"),
    ("installed_at", "Установлена"),
    ("warranty_until", "Гарантия до"),
    ("last_service_at", "Последнее ТО"),
    ("service_interval_months", "Интервал ТО, мес."),
    ("notes", "Примечания"),
)

MISSING = "НЕТ"


def _plain(value: Any) -> Any:
    """Приводит значение к виду, пригодному и для Excel, и для CSV.

    Даты оставляем объектами date: Excel покажет их как настоящие даты,
    по которым можно сортировать и фильтровать. CSV-писатель всё равно
    превратит их в ГГГГ-ММ-ДД.
    """
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, Decimal):
        return float(value)
    return "" if value is None else value


def render_csv(table: Table) -> bytes:
    """CSV с точкой с запятой и BOM — так его открывает Excel.

    BOM нужен, потому что иначе Excel на русской локали читает файл
    в однобайтовой кодировке и показывает кракозябры.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_ALL)
    writer.writerow(table.headers)
    for row in table.rows:
        writer.writerow(
            [v.isoformat() if isinstance(v, (date, datetime)) else v for v in row]
        )
    return "﻿".encode() + buffer.getvalue().encode("utf-8")


def render_xlsx(tables: list[Table]) -> bytes:
    """Книга Excel: по листу на таблицу, шапка закреплена и выделена.

    openpyxl импортируем внутри функции, а не сверху файла: так модуль
    грузится только когда действительно просят Excel, и приложение
    поднимется даже если библиотеку забыли поставить.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)  # пустой лист, который создаётся по умолчанию

    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EDEDED")
    bad_fill = PatternFill("solid", fgColor="FFC7CE")

    for table in tables:
        ws = wb.create_sheet(title=table.name[:31])  # Excel режет на 31 символе
        ws.append(table.headers)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        for row in table.rows:
            ws.append(row)
            # Подсветка дыр в комплекте: глазами такую таблицу
            # иначе не прочитать.
            for cell in ws[ws.max_row]:
                if cell.value == MISSING or (
                    isinstance(cell.value, str) and cell.value.endswith("(истёк)")
                ):
                    cell.fill = bad_fill

        # Ширина по самому длинному значению, но не уже 9 и не шире 42:
        # иначе «Примечания» растянут лист на пол-экрана.
        for idx, header in enumerate(table.headers, start=1):
            longest = max(
                [len(str(header))]
                + [len(str(r[idx - 1])) for r in table.rows if r[idx - 1] != ""],
                default=10,
            )
            ws.column_dimensions[get_column_letter(idx)].width = min(
                max(longest + 2, 9), 42
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class ExportService:
    def __init__(self, session: AsyncSession) -> None:
        # list_with_relations, а не list: у list есть limit со значением
        # по умолчанию, и на сто первой площадке выгрузка молча
        # обрезалась бы — такое потом ищут неделями.
        self.stations = StationRepository(session)

    async def stations_table(self) -> Table:
        rows = []
        for st in await self.stations.list_with_relations():
            row = []
            for field, _ in STATION_COLUMNS:
                value = getattr(st, field)
                if field == "status":
                    value = STATION_STATUS_RU.get(value, value)
                row.append(_plain(value))
            row.append(len(st.chargers))
            row.append(sum(1 for d in st.documents if d.is_current))
            rows.append(row)

        headers = [title for _, title in STATION_COLUMNS]
        headers += ["Зарядных станций", "Документов"]
        return Table("Площадки", headers, rows)

    async def chargers_table(self) -> Table:
        rows = []
        for st in await self.stations.list_with_relations():
            for ch in st.chargers:
                row: list[Any] = [st.code or "", st.name]
                for field, _ in CHARGER_COLUMNS:
                    value = getattr(ch, field)
                    if field == "status":
                        value = CHARGER_STATUS_RU.get(value, value)
                    row.append(_plain(value))
                rows.append(row)

        headers = ["Код площадки", "Площадка"]
        headers += [title for _, title in CHARGER_COLUMNS]
        return Table("Оборудование", headers, rows)

    async def compliance_table(self, today: date | None = None) -> Table:
        """Матрица комплектности: площадки по строкам, типы документов
        по столбцам.

        В клетке либо «НЕТ» (обязательного документа не хватает), либо
        дата окончания срока, либо «есть» для бессрочных. Такой лист
        можно распечатать и принести на совещание.
        """
        today = today or date.today()
        doc_codes = [d.code for d in docs.DOC_TYPES]

        rows = []
        for st in await self.stations.list_with_relations():
            facts = station_to_facts(st, list(st.documents), list(st.chargers))
            result = rules.evaluate_station(facts, today)
            by_code = {row.code: row for row in result.checklist}

            row: list[Any] = [
                st.code or "",
                st.name,
                STATION_STATUS_RU.get(st.status, st.status),
                sum(1 for a in result.alerts if a.level == AlertLevel.OVERDUE),
                sum(1 for a in result.alerts if a.level == AlertLevel.WARNING),
            ]
            for code in doc_codes:
                cell = by_code.get(code)
                if cell is None or cell.count == 0:
                    # Обязательный и отсутствует — дырка в комплекте.
                    # Необязательный просто оставляем пустым.
                    row.append(MISSING if (cell and cell.required) else "")
                elif cell.valid_until:
                    overdue = cell.days is not None and cell.days < 0
                    stamp = cell.valid_until.isoformat()
                    row.append(f"{stamp} (истёк)" if overdue else stamp)
                else:
                    row.append("есть")
            rows.append(row)

        headers = ["Код", "Площадка", "Стадия", "Просрочено", "Истекает"]
        headers += [docs.doc_title(code) for code in doc_codes]
        return Table("Комплектность", headers, rows)

    async def workbook(self) -> bytes:
        """Вся сеть одной книгой: три листа вместо трёх файлов."""
        return render_xlsx(
            [
                await self.stations_table(),
                await self.chargers_table(),
                await self.compliance_table(),
            ]
        )

    async def stations_json(self) -> dict:
        """Полная выгрузка со вложенными списками — для переноса данных
        и разбора сторонними инструментами, а не для чтения человеком."""
        payload = []
        for st in await self.stations.list_with_relations():
            item = {
                col.key: _json_safe(getattr(st, col.key))
                for col in st.__table__.columns
            }
            item["chargers"] = [
                {c.key: _json_safe(getattr(ch, c.key)) for c in ch.__table__.columns}
                for ch in st.chargers
            ]
            item["documents"] = [
                {
                    "id": d.id,
                    "doc_type": d.doc_type,
                    "doc_type_title": docs.doc_title(d.doc_type),
                    "title": d.title,
                    "doc_number": d.doc_number,
                    "doc_date": _json_safe(d.doc_date),
                    "valid_until": _json_safe(d.valid_until),
                    "version": d.version,
                    "is_current": d.is_current,
                    "original_name": d.original_name,
                }
                for d in st.documents
            ]
            payload.append(item)

        return {
            "exported_at": datetime.now().isoformat(timespec="seconds"),
            "count": len(payload),
            "stations": payload,
        }


def _json_safe(value: Any) -> Any:
    """Для JSON булево оставляем булевым, а дату отдаём строкой."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value
