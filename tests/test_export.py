"""Выгрузка: что именно попадает в файл.

Проверяем отрисовку, а не сбор данных из базы: таблицы собираем руками.
"""
import io
from datetime import date

from app.service.export import MISSING, Table, render_csv, render_xlsx


def sample() -> Table:
    return Table(
        name="Площадки",
        headers=["Код", "Название", "Открыта", "Навес"],
        rows=[
            ["ASKO-01", "Махачкала, Центр", date(2024, 3, 1), "да"],
            ["ASKO-05", 'ТД «Киргу»', "", "нет"],
        ],
    )


# ── CSV ──────────────────────────────────────────────────────────────────────
def test_csv_начинается_с_метки_кодировки():
    """Без неё Excel на русской локали покажет кракозябры."""
    assert render_csv(sample()).startswith(b"\xef\xbb\xbf")


def test_csv_разделяется_точкой_с_запятой():
    """Запятая не годится: в русской локали Excel ждёт точку с запятой,
    а с запятой складывает всю строку в одну ячейку."""
    text = render_csv(sample()).decode("utf-8-sig")
    assert text.splitlines()[0] == '"Код";"Название";"Открыта";"Навес"'


def test_csv_сохраняет_кириллицу():
    text = render_csv(sample()).decode("utf-8-sig")
    assert "Махачкала, Центр" in text
    assert "ТД «Киргу»" in text


def test_csv_пишет_даты_в_читаемом_виде():
    text = render_csv(sample()).decode("utf-8-sig")
    assert "2024-03-01" in text


def test_csv_экранирует_разделитель_внутри_значения():
    """В названии площадки есть запятая, а в примечаниях может
    оказаться и точка с запятой — строка не должна разъехаться."""
    table = Table("т", ["Название"], [['Манас; ФАД "Кавказ", 842 км']])
    import csv as csv_mod

    text = render_csv(table).decode("utf-8-sig")
    rows = list(csv_mod.reader(io.StringIO(text), delimiter=";"))
    assert rows[1] == ['Манас; ФАД "Кавказ", 842 км']


# ── Excel ────────────────────────────────────────────────────────────────────
def test_xlsx_это_настоящий_файл_excel():
    """Формат xlsx — ZIP-архив, он начинается с сигнатуры PK."""
    assert render_xlsx([sample()]).startswith(b"PK")


def test_xlsx_раскладывает_таблицы_по_листам():
    from openpyxl import load_workbook

    data = render_xlsx([sample(), Table("Оборудование", ["Серийный"], [["SN1"]])])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Площадки", "Оборудование"]


def test_xlsx_закрепляет_шапку_и_включает_фильтр():
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(render_xlsx([sample()])))["Площадки"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold is True


def test_xlsx_хранит_даты_датами():
    """Если записать дату строкой, по ней нельзя будет отсортировать."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(render_xlsx([sample()])))["Площадки"]
    assert ws["C2"].value.date() == date(2024, 3, 1)


def test_xlsx_подсвечивает_дыры_в_комплекте():
    from openpyxl import load_workbook

    table = Table(
        "Комплектность",
        ["Код", "Аренда", "Проект"],
        [["ASKO-01", MISSING, "2026-09-01 (истёк)"], ["ASKO-05", "есть", "есть"]],
    )
    ws = load_workbook(io.BytesIO(render_xlsx([table])))["Комплектность"]

    painted = {
        c.value
        for row in ws.iter_rows(min_row=2)
        for c in row
        if c.fill.fgColor.rgb == "00FFC7CE"
    }
    assert painted == {MISSING, "2026-09-01 (истёк)"}


def test_xlsx_не_растягивает_колонку_на_весь_экран():
    """Длинное примечание не должно делать лист нечитаемым."""
    from openpyxl import load_workbook

    long_text = "очень длинное примечание " * 20
    ws = load_workbook(
        io.BytesIO(render_xlsx([Table("т", ["Примечания"], [[long_text]])]))
    )["т"]
    assert ws.column_dimensions["A"].width <= 42
